#!/usr/bin/env python3
"""Bulk-ingest CMI "Save as Excel / Export Page Data" XLSX files into the five-year snapshot.

Fallback path when browser automation cannot reliably switch/export every page.
The script never overwrites a conflicting existing observation silently.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode

from openpyxl import load_workbook

CATALOG = Path("data/cmi/catalog/indicators.json")
MERGED = Path("output/cmi_5y/health_kpi_records.json")
NORMALIZED = Path("data/cmi/normalized")
MANIFESTS = Path("data/cmi/manifests")
YEARS = [2565,2566,2567,2568,2569]
BASE = "https://cmi.maewanghospital.go.th/web/index.php"
HOSP_RE = re.compile(r"^\s*(\d{5})\s*(.*)$")
YEAR_RE = re.compile(r"\b(25\d{2})\b")


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def clean(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = clean(v).replace(",", "").replace("%", "")
    if t in {"", "-", "—"}:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", t)
    return float(m.group(0)) if m else None


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def source_url(code, family):
    if family == "service_plan":
        return f"{BASE}?{urlencode({'co_thip_new':code,'r':'service/index'})}"
    return f"{BASE}?{urlencode({'id':code,'r':'report/drgindexreport'})}"


def header_index(header, terms):
    for i, cell in enumerate(header):
        c = clean(cell).lower()
        if any(term.lower() in c for term in terms):
            return i
    return None


def detect_unit(text):
    t = clean(text).lower()
    if "ต่อแสน" in t or "100000" in t or "100,000" in t:
        return "/100k"
    if "ต่อพัน" in t or "1000" in t or "1,000" in t:
        return "/1000"
    if "ร้อยละ" in t or "%" in t:
        return "%"
    if "adjrw" in t:
        return "AdjRW"
    return ""


def detect_code_and_year(matrix, filename, catalog):
    head = " ".join([filename] + [clean(x) for row in matrix[:40] for x in row if x is not None])
    upper = head.upper()
    codes = sorted(catalog, key=len, reverse=True)
    code = next((c for c in codes if re.search(rf"(^|[^A-Z0-9]){re.escape(c)}([^A-Z0-9]|$)", upper)), None)
    ym = YEAR_RE.search(head)
    year = int(ym.group(1)) if ym else None
    return code, year


def parse_matrix(matrix, file_path, code, year, indicator):
    header_row = None
    header_idx = None
    for i, row in enumerate(matrix[:30]):
        joined = " | ".join(clean(x) for x in row)
        if "สถานพยาบาล" in joined or "หน่วยบริการ" in joined:
            header_row = row
            header_idx = i
            break
    if header_row is None:
        return []

    hosp_col = header_index(header_row, ["สถานพยาบาล","หน่วยบริการ"])
    province_col = header_index(header_row, ["จังหวัด"])
    level_col = header_index(header_row, ["ระดับ"])
    numerator_col = header_index(header_row, ["ตัวตั้ง","numerator"])
    denominator_col = header_index(header_row, ["ตัวหาร","denominator"])
    value_col = header_index(header_row, ["ร้อยละ","อัตรา","ผลลัพธ์","ค่าเฉลี่ย","ค่า"])
    unit = detect_unit(" ".join(clean(x) for x in header_row))
    file_hash = sha256_file(file_path)
    records = []

    for row in matrix[header_idx + 1:]:
        if not any(clean(x) for x in row):
            continue
        indexes = [hosp_col] if hosp_col is not None and hosp_col < len(row) else range(len(row))
        hi = None
        hm = None
        for idx in indexes:
            m = HOSP_RE.match(clean(row[idx]))
            if m:
                hi, hm = idx, m
                break
        if hi is None:
            continue

        province = clean(row[province_col]) if province_col is not None and province_col < len(row) else ""
        level = clean(row[level_col]) if level_col is not None and level_col < len(row) else (clean(row[hi-1]) if hi > 0 else "")
        numerator = num(row[numerator_col]) if numerator_col is not None and numerator_col < len(row) else None
        denominator = num(row[denominator_col]) if denominator_col is not None and denominator_col < len(row) else None
        value = num(row[value_col]) if value_col is not None and value_col < len(row) else None

        tail = [num(x) for x in row[hi+1:]]
        tail = [x for x in tail if x is not None]
        if numerator is None and denominator is None and len(tail) >= 3:
            numerator, denominator = tail[0], tail[1]
        if value is None and tail:
            value = tail[-1]
        if value is None and numerator is None and denominator is None:
            continue

        records.append({
            "year": year,
            "indicator_code": code,
            "indicator_name": indicator.get("name") or code,
            "province": province,
            "hospital_code": hm.group(1),
            "hospital_name": clean(hm.group(2)),
            "service_level": level,
            "numerator": numerator,
            "denominator": denominator,
            "value": value,
            "unit": unit,
            "definition_text": "",
            "definition_hash": None,
            "source_url": indicator.get("source_url") or source_url(code, indicator.get("family")),
            "source_file": str(file_path),
            "source_sha256": file_hash,
            "collected_at": now_iso(),
            "source_type": "cmi_excel_export",
        })
    return records


def load_matrix(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    matrices = []
    for ws in wb.worksheets:
        matrix = [list(row) for row in ws.iter_rows(values_only=True)]
        matrices.append((ws.title, matrix))
    return matrices


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_dir", help="directory containing CMI .xlsx exports")
    args = ap.parse_args()

    catalog_doc = json.loads(CATALOG.read_text(encoding="utf-8"))
    indicator_map = {str(x["code"]).upper(): x for x in catalog_doc.get("indicators", [])}
    existing_doc = json.loads(MERGED.read_text(encoding="utf-8")) if MERGED.exists() else {
        "schema_version":"nco-cmi-kpi-5y-v1","years":YEARS,"records":[]
    }
    existing = {(r["year"],r["indicator_code"],r["hospital_code"]): r for r in existing_doc.get("records", [])}
    conflicts = []
    inserted = 0
    skipped_files = []

    for path in sorted(Path(args.input_dir).rglob("*.xlsx")):
        parsed_any = False
        try:
            for sheet_name, matrix in load_matrix(path):
                code, year = detect_code_and_year(matrix, path.name + " " + sheet_name, indicator_map.keys())
                if not code or year not in YEARS:
                    continue
                rows = parse_matrix(matrix, path, code, year, indicator_map[code])
                if not rows:
                    continue
                parsed_any = True
                for row in rows:
                    key = (row["year"],row["indicator_code"],row["hospital_code"])
                    prior = existing.get(key)
                    if not prior:
                        existing[key] = row
                        inserted += 1
                        continue
                    comparable = ("numerator","denominator","value")
                    same = all(prior.get(k) == row.get(k) for k in comparable)
                    if not same:
                        conflicts.append({
                            "key":{"year":key[0],"indicator_code":key[1],"hospital_code":key[2]},
                            "existing":{k:prior.get(k) for k in comparable},
                            "incoming":{k:row.get(k) for k in comparable},
                            "incoming_file":str(path),
                        })
            if not parsed_any:
                skipped_files.append({"file":str(path),"reason":"could not detect indicator/year/table"})
        except Exception as exc:
            skipped_files.append({"file":str(path),"reason":str(exc)})

    records = sorted(existing.values(), key=lambda r:(r["year"],r["indicator_code"],r.get("province",""),r["hospital_code"]))
    indicator_years = {(r["indicator_code"],r["year"]) for r in records}
    expected = len(indicator_map) * len(YEARS)
    parsed = len(indicator_years)
    existing_doc.update({
        "schema_version":"nco-cmi-kpi-5y-v1",
        "generated_at":now_iso(),
        "source_system":"CMI / Service Plan Region 1",
        "source_base_url":BASE,
        "years":YEARS,
        "status":"incomplete" if conflicts or parsed < expected else "complete",
        "records":records,
        "catalog_size":len(indicator_map),
        "completeness":{"expected":expected,"parsed":parsed,"percent":round(parsed/expected*100,2) if expected else 0},
    })
    MERGED.parent.mkdir(parents=True, exist_ok=True)
    MERGED.write_text(json.dumps(existing_doc, ensure_ascii=False, indent=2)+"\n",encoding="utf-8")

    NORMALIZED.mkdir(parents=True, exist_ok=True)
    for year in YEARS:
        year_rows=[r for r in records if r["year"]==year]
        (NORMALIZED/f"{year}.json").write_text(json.dumps({
            "schema_version":"nco-cmi-normalized-v1","year":year,"records":year_rows
        },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")

    MANIFESTS.mkdir(parents=True, exist_ok=True)
    (MANIFESTS/"import_conflicts.json").write_text(json.dumps({
        "schema_version":"nco-cmi-import-conflicts-v1",
        "generated_at":now_iso(),
        "conflicts":conflicts,
        "skipped_files":skipped_files,
    },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")

    csv_path=MERGED.with_suffix(".csv")
    fields=["year","indicator_code","indicator_name","province","hospital_code","hospital_name","service_level",
            "numerator","denominator","value","unit","source_url","source_sha256","source_type","collected_at"]
    with csv_path.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for row in records: w.writerow({k:row.get(k) for k in fields})

    print(f"Inserted {inserted} observations; conflicts={len(conflicts)}; skipped_files={len(skipped_files)}")
    print(f"Indicator-year completeness: {parsed}/{expected}")
    return 1 if conflicts else 0


if __name__ == "__main__":
    raise SystemExit(main())
